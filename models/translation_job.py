# -*- coding: utf-8 -*-
import os
import io
import json
import logging
import threading
import base64
import sys
import time
from typing import List, Dict, Any, Tuple, Optional

# Odoo imports
from odoo import models, fields, api, _
from odoo.exceptions import UserError
from odoo.tools import config

# OCR & PDF Imports
import fitz  # PyMuPDF
from PIL import Image
try:
    import pytesseract
except OSError:
    pass

# Word Compilation Imports
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Fallback Translator Imports
from deep_translator import GoogleTranslator

_logger = logging.getLogger(__name__)

# Interceptor de consola para evitar UnicodeEncodeError en Windows
if sys.platform.startswith('win'):
    class SafeUnicodeStdout:
        def __init__(self, original_stdout):
            self.original_stdout = original_stdout
        def write(self, data):
            try: 
                self.original_stdout.write(data)
            except UnicodeEncodeError: 
                self.original_stdout.write(data.encode('ascii', errors='replace').decode('ascii'))
        def __getattr__(self, name): 
            return getattr(self.original_stdout, name)
    sys.stdout = SafeUnicodeStdout(sys.stdout)

# =====================================================================
# MAIN ODOO MODEL
# =====================================================================

class TranslationJob(models.Model):
    _name = 'translation.job'
    _description = 'Trabajo de Traducción Multi-Agente Híbrido'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(string='Referencia', required=True, copy=False, readonly=True, default=lambda self: _('Nuevo'))
    state = fields.Selection([
        ('draft', 'Borrador'),
        ('processing', 'Procesando OCR/Agentes'),
        ('alert', 'Conflicto Narrativo'),
        ('done', 'Finalizado'),
        ('error', 'Error')
    ], string='Estado', default='draft', tracking=True)

    # NUEVO CAMPO: Selector Comercial de País
    target_country = fields.Selection([
        ('US', 'Estados Unidos'),
        ('DE', 'Alemania'),
        ('JP', 'Japón'),
        ('UK', 'Reino Unido'),
        ('CA', 'Canadá'),
        ('CN', 'China'),
        ('FR', 'Francia'),
        ('IT', 'Italia'),
        ('ES', 'España'),
        ('NL', 'Países Bajos')
    ], string='País Destino (Contexto Comercial)', default='ES', required=True, tracking=True)

    # ACTUALIZADO: Idiomas expandidos
    source_lang = fields.Selection([
        ('auto', 'Detectar automáticamente'),
        ('en_US', 'Inglés (EE.UU.)'),
        ('en_GB', 'Inglés (Reino Unido)'),
        ('es_ES', 'Español (España)'),
        ('es_419', 'Español (Latam)'),
        ('pt', 'Portugués'),
        ('fr', 'Francés'),
        ('de', 'Alemán'),
        ('it', 'Italiano'),
        ('ja', 'Japonés'),
        ('zh', 'Chino Mandarín'),
        ('nl', 'Neerlandés')
    ], string='Idioma Origen', default='auto', required=True)

    # ACTUALIZADO: Idiomas expandidos
    target_lang = fields.Selection([
        ('es_ES', 'Español (España)'),
        ('es_419', 'Español (Latam)'),
        ('en_US', 'Inglés (EE.UU.)'),
        ('en_GB', 'Inglés (Reino Unido)'),
        ('pt', 'Portugués'),
        ('fr', 'Francés'),
        ('de', 'Alemán'),
        ('it', 'Italiano'),
        ('ja', 'Japonés'),
        ('zh', 'Chino Mandarín'),
        ('nl', 'Neerlandés')
    ], string='Idioma Destino', default='es_ES', required=True)

    pdf_file = fields.Binary(string='Archivo PDF Original', required=True)
    pdf_filename = fields.Char(string='Nombre PDF')
    docx_file = fields.Binary(string='Documento Word Traducido', readonly=True)
    docx_filename = fields.Char(string='Nombre Word')
    line_ids = fields.One2many('translation.job.line', 'job_id', string='Segmentos Traducidos')

    # Campos de soporte para la gestión de conflictos y ventana modal
    conflict_text_original = fields.Text(string='Texto del Conflicto Original')
    conflict_suggested_options = fields.Text(string='Opciones de Traducción Sugeridas (JSON)')
    current_page_num = fields.Integer(string='Página del Conflicto')
    current_block_num = fields.Integer(string='Bloque del Conflicto')
    current_block_alignment = fields.Selection([
        ('left', 'Izquierda'),
        ('center', 'Centro'),
        ('right', 'Derecha')
    ], string='Alineación del Conflicto', default='left')
    current_block_is_heading = fields.Boolean(string='Es Título/Encabezado en Conflicto', default=False)

    # Lock global en memoria y set para evitar ejecuciones concurrentes de hilos en el mismo job_id
    _active_jobs = set()
    _active_jobs_lock = threading.Lock()

    @api.model_create_multi
    def create(self, vals_list):
        """Sobreescritura correcta del método create en lote utilizando model_create_multi"""
        for vals in vals_list:
            if vals.get('name', _('Nuevo')) == _('Nuevo'):
                vals['name'] = self.env['ir.sequence'].next_by_code('translation.job') or _('Nuevo')
        return super(TranslationJob, self).create(vals_list)

    def _is_test_mode(self):
        """Comprueba de forma robusta si estamos ejecutando bajo el framework de pruebas de Odoo."""
        return (
            config.get('test_enable') or 
            self.env.registry.in_test_mode() or 
            getattr(threading.current_thread(), 'testing', False)
        )

    def _safe_commit(self):
        """Realiza commit de base de datos solo si no nos encontramos en entorno de pruebas/unittests."""
        if not self._is_test_mode():
            self.env.cr.commit()

    def _sanitize_original_text(self, text):
        """Limpia caracteres de control, comillas sin escapar y espacios duplicados."""
        if not text:
            return ""
        import re
        # Reemplazar caracteres de control por espacios
        text = re.sub(r'[\x00-\x1F\x7F-\x9F]', ' ', text)
        # Eliminar comillas dobles y simples no escapadas
        text = re.sub(r'(?<!\\)"', '', text)
        text = re.sub(r"(?<!\\)'", "", text)
        # Reemplazar múltiples espacios por uno solo
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    def _extract_pages_from_attachment(self):
        """
        Extrae el texto del PDF de forma híbrida:
        1. Intenta extracción nativa rápida con PyMuPDF (fitz).
        2. Si la página está vacía, tiene menos de 10 caracteres (ruido/imágenes) o falla,
           realiza un fallback automático a Tesseract OCR (pytesseract).
        Garantiza la liberación inmediata de buffers de imagen para control de memoria RAM.
        Retorna: List[Dict[str, Any]] -> [{'page_num': 0, 'text': '...'}, ...]
        """
        self.ensure_one()
        
        # Configuración por defecto de la ruta de Tesseract en Windows
        if sys.platform.startswith('win'):
            tesseract_default_path = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
            if os.path.exists(tesseract_default_path):
                pytesseract.pytesseract.tesseract_cmd = tesseract_default_path
                _logger.info("Ruta de Tesseract configurada en Windows: %s", tesseract_default_path)
            else:
                _logger.warning("No se encontró el ejecutable de Tesseract en la ruta por defecto de Windows: %s", tesseract_default_path)

        pages_data = []
        try:
            pdf_data = base64.b64decode(self.pdf_file)
            doc_fitz = fitz.open(stream=pdf_data, filetype="pdf")
            page_count = len(doc_fitz)
            _logger.info("PDF cargado. Total de páginas detectadas: %s", page_count)
            
            for page_num in range(page_count):
                page = doc_fitz[page_num]
                page_label = f"Página {page_num + 1}/{page_count}"
                
                # 1. Intentar extracción nativa rápida
                text_content = ""
                try:
                    text_content = page.get_text("text").strip()
                except Exception as ext_err:
                    _logger.warning("[%s] Error al extraer texto nativo: %s", page_label, ext_err)

                # 2. Lógica de decisión dinámica (vacío, menos de 10 caracteres o falla)
                if not text_content or len(text_content) < 10:
                    _logger.info("[%s] Detectada página escaneada/imagen (Largo nativo: %s). Iniciando Tesseract OCR...", 
                                 page_label, len(text_content) if text_content else 0)
                    
                    try:
                        # Renderizar la página a pixmap (imagen en memoria a 200 DPI para precisión de OCR)
                        pix = page.get_pixmap(dpi=200)
                        img_bytes = pix.tobytes("png")
                        
                        # Carga de imagen en PIL para pasar a Tesseract
                        with Image.open(io.BytesIO(img_bytes)) as img:
                            # Intentamos inglés + español para mayor fidelidad
                            try:
                                text_content = pytesseract.image_to_string(img, lang='eng+spa').strip()
                            except Exception:
                                # Fallback a idioma por defecto (ej. inglés)
                                text_content = pytesseract.image_to_string(img).strip()
                        
                        # Limpieza explícita de objetos de memoria
                        del pix
                        del img_bytes
                        
                        if text_content:
                            _logger.info("[%s] Texto extraído vía OCR exitosamente (Largo: %s).", page_label, len(text_content))
                        else:
                            text_content = "[Página vacía]"
                            _logger.info("[%s] El OCR no detectó caracteres. Página marcada como vacía.", page_label)
                            
                    except Exception as ocr_err:
                        _logger.error("[%s] Fallo crítico durante la ejecución de Tesseract OCR: %s", page_label, ocr_err)
                        text_content = "[Error en OCR - Imagen no legible]"
                else:
                    _logger.info("[%s] Texto nativo extraído (Largo: %s).", page_label, len(text_content))
                
                pages_data.append({
                    'page_num': page_num,
                    'text': text_content
                })
                
            return pages_data
        except Exception as e:
            _logger.exception("Error crítico no recuperable en _extract_pages_from_attachment:")
            raise UserError(_("Fallo en la extracción de páginas del PDF: %s") % str(e))

    def action_start_translation(self):
        """Inicia el proceso de traducción asíncrono y levanta el hilo secundario con un cursor seguro."""
        self.ensure_one()
        # Si ya está procesando, reanudamos sin borrar segmentos existentes ni lanzar error.
        if self.state == 'processing':
            if self._is_test_mode():
                return {'type': 'ir.actions.client', 'tag': 'reload'}
        else:
            self.write({'state': 'processing', 'line_ids': [(5, 0, 0)]})
            self._safe_commit()
        
        # En modo unit-test ejecutamos de forma directa síncrona sin crear hilos paralelos ni nuevos cursores
        if self._is_test_mode():
            self._process_translation_background()
            return {'type': 'ir.actions.client', 'tag': 'reload'}
        
        # Apertura de cursor independiente
        new_cr = self.pool.cursor()
        new_env = api.Environment(new_cr, self.env.uid, self.env.context)
        job_in_thread = new_env['translation.job'].browse(self.id)
        
        # Lanzamiento asíncrono del hilo secundario
        threading.Thread(target=job_in_thread._process_translation_background, daemon=True).start()
        return {'type': 'ir.actions.client', 'tag': 'reload'}

    def _process_translation_background(self):
        """Procesa y traduce el PDF en segundo plano con control de errores atómico por página."""
        # Control estricto de concurrencia: Evitar ejecutar hilos duplicados para el mismo registro
        with self._active_jobs_lock:
            if self.id in self._active_jobs:
                _logger.warning("El trabajo ID %s ya tiene un hilo secundario activo. Abortando ejecución duplicada.", self.id)
                return
            self._active_jobs.add(self.id)

        # Extraer nombres humanos de los idiomas para el prompt
        source_lang_name = dict(self._fields['source_lang'].selection).get(self.source_lang, "Detectar automáticamente")
        target_lang_name = dict(self._fields['target_lang'].selection).get(self.target_lang, "Español (España)")

        # Enrutador Lingüístico: Obtener el enfoque sectorial basado en el país destino
        country_code = self.target_country
        enfoque_sectorial = ""
        if country_code == 'US':
            enfoque_sectorial = "Prioriza un tono corporativo multinacional. Aplica estrictamente normativas de acceso lingüístico para traducciones de índole médica, judicial o gubernamental."
        elif country_code == 'DE':
            enfoque_sectorial = "Fuerza un enfoque de traducción de alta precisión técnica, adaptado a manuales de ingeniería, automoción y manufactura pesada."
        elif country_code == 'JP':
            enfoque_sectorial = "Ajusta los algoritmos para localización de contenidos de entretenimiento, software y exportación tecnológica, manteniendo estrictamente el protocolo de cortesía cultural (Keigo)."
        elif country_code == 'UK':
            enfoque_sectorial = "Configura el tono para alta precisión financiera, corporativa y del sector legal/jurídico internacional."
        elif country_code == 'CA':
            enfoque_sectorial = "Activa el protocolo de bilingüismo oficial. Valida que la documentación, etiquetado y servicios públicos cumplan estrictamente las normativas equivalentes de inglés y francés canadiense."
        elif country_code == 'CN':
            enfoque_sectorial = "Optimiza la traducción bidireccional hacia el chino mandarín comercial, enfocada en plataformas de comercio electrónico y expansión global."
        elif country_code == 'FR':
            enfoque_sectorial = "Adapta la salida para localización de medios, doblaje, subtitulado y cumplimiento estricto de las regulaciones comerciales de protección idiomática de la UE."
        elif country_code == 'IT':
            enfoque_sectorial = "Orienta el estilo estilístico hacia los sectores de la moda, diseño de gama alta, turismo y exportación de maquinaria industrial."
        elif country_code == 'ES':
            enfoque_sectorial = "Prioriza la adaptación de contenidos digitales, flujos turísticos y la diferenciación semántica si el contenido requiere puente directo con el español de Latinoamérica."
        elif country_code == 'NL':
            enfoque_sectorial = "Enfoca la estructura sintáctica hacia la logística, el e-commerce transfronterizo y flujos corporativos de entrada europea."

        _logger.info("Iniciando procesamiento de traducción de fondo para el trabajo ID %s", self.id)
        try:
            # 1. Configurar la API de Gemini
            api_key = self.env['ir.config_parameter'].sudo().get_param('translation_agent_intel.gemini_api_key')
            model = None
            if api_key:
                try:
                    import google.generativeai as genai
                    genai.configure(api_key=api_key)
                    model = genai.GenerativeModel("gemini-1.5-flash")
                    _logger.info("API de Gemini configurada correctamente.")
                except Exception as genai_err:
                    _logger.warning(f"No se pudo inicializar la API de Gemini: {genai_err}")

            # Extracción híbrida inteligente del contenido de páginas
            extracted_pages = self._extract_pages_from_attachment()
            
            # Cargar el documento PDF para operaciones de renderizado/pixmap si fuese necesario
            try:
                pdf_data = base64.b64decode(self.pdf_file)
                doc_fitz = fitz.open(stream=pdf_data, filetype="pdf")
            except Exception as pdf_err:
                _logger.exception("Error al abrir o decodificar el archivo PDF para procesamiento secundario:")
                raise pdf_err

            # 3. Iteración sobre las páginas del documento
            for page_info in extracted_pages:
                page_num = page_info['page_num']
                page_text = page_info['text']
                _logger.info("Procesando página %s/%s", page_num + 1, len(extracted_pages))
                try:
                    # Usamos savepoint para evitar invalidar la transacción completa por fallas de base de datos en una página
                    with self.env.cr.savepoint():
                        blocks_to_translate = []
                        if page_text not in ["[Página vacía]", "[Error en OCR - Imagen no legible]"]:
                            # Separar por saltos de párrafo dobles (funciona idéntico para texto nativo y OCR)
                            raw_paras = page_text.split("\n\n")
                            native_paragraphs = []
                            for rp in raw_paras:
                                # Limpiar saltos de línea internos dentro de cada párrafo
                                cleaned_para = " ".join(line.strip() for line in rp.split("\n") if line.strip())
                                if cleaned_para:
                                    native_paragraphs.append(cleaned_para)
                            
                            _logger.info("Página %s: Detectados %s párrafos para traducir.", page_num + 1, len(native_paragraphs))
                            seq = 0
                            for text in native_paragraphs:
                                sanitized = self._sanitize_original_text(text)
                                # Evitamos fragmentos vacíos, números puros o strings demasiado cortos
                                if sanitized and not sanitized.isdigit() and len(sanitized) > 2:
                                    # Heurística para títulos (mayúsculas y longitud corta)
                                    is_heading = sanitized.isupper() and len(sanitized) < 120
                                    blocks_to_translate.append({
                                        "seq": seq,
                                        "text": sanitized,
                                        "is_heading": is_heading,
                                        "alignment": "left"
                                    })
                                    seq += 1
                        
                        if not blocks_to_translate:
                            _logger.info("Página %s: Sin contenido para traducir.", page_num + 1)
                            continue

                        # Arquitectura Multi-Agente (Prompt Chaining) por Lote para la página actual
                        translations_map = {}
                        if model:
                            try:
                                # Preparar payload
                                clean_payload = [{"seq": b["seq"], "text": b["text"]} for b in blocks_to_translate]
                                payload_json = json.dumps(clean_payload, ensure_ascii=False)
                                
                                # NODO 1: Sanitización
                                prompt_nodo1 = f"""Actúas como un Ingeniero de Datos experto en limpieza de texto y corrección de OCR. Analiza el contenido inyectado en [TEXTO_A_TRADUCIR] en formato JSON y realiza las siguientes tareas de pre-procesamiento:
- Detecta y elimina por completo cualquier cadena de texto corrupta, símbolos extraños o ruido de escaneo (ej. secuencias tipo 'H UW G KH...').
- Une de forma lógica las frases que hayan quedado cortadas a mitad de palabra por culpa de saltos de página.
- Si detectas que la línea temporal de la narrativa se corta abruptamente y se repite o continúa más abajo, reordena los párrafos para que mantengan una coherencia cronológica perfecta de principio a fin.
Output: Devuelve obligatoriamente un JSON válido manteniendo la estructura (seq, text) con el texto original limpio y ordenado, sin comentarios.
[TEXTO_A_TRADUCIR]: {payload_json}"""
                                
                                _logger.info("Página %s: Ejecutando NODO 1 (Sanitización)...", page_num + 1)
                                resp1 = model.generate_content(prompt_nodo1, generation_config={"response_mime_type": "application/json"})
                                json_nodo1 = resp1.text.replace("```json", "").replace("```", "").strip()
                                
                                # NODO 2: Traducción
                                prompt_nodo2 = f"""Actúas como un Traductor Literario y Localizador Profesional especializado en la combinación de idiomas desde {source_lang_name} hacia {target_lang_name}.
Traduce el texto procesado por el Nodo 1 (en formato JSON) aplicando estrictamente estas reglas de negocio:
- Contexto de Mercado (País Destino {country_code}): {enfoque_sectorial}
- Prohibida la traducción literal: Identifica modismos, expresiones antiguas, argot y metáforas culturales. Busca y aplica su equivalente semántico exacto y natural en el idioma de destino (ej. 'sugar-hogshead' debe localizarse como 'barril de azúcar' y 'stretchers' como 'exageraciones/mentiras piadosas').
- Consistencia: Asegúrate de traducir absolutamente todo el texto. Bajo ninguna circunstancia dejes fragmentos o frases completas en el idioma original si el resto del documento ha sido procesado.
- Fluidez Nativa: La estructura sintáctica final debe sonar natural para un hablante nativo del idioma de destino, evitando calcos gramaticales del idioma origen.
Devuelve únicamente un JSON válido manteniendo la estructura (seq, text).
Texto a procesar: {json_nodo1}"""
                                
                                _logger.info("Página %s: Ejecutando NODO 2 (Traducción)...", page_num + 1)
                                resp2 = model.generate_content(prompt_nodo2, generation_config={"response_mime_type": "application/json"})
                                json_nodo2 = resp2.text.replace("```json", "").replace("```", "").strip()

                                # NODO 3: Auditoría y QA
                                prompt_nodo3 = f"""Actúas como un Corrector de Estilo y Auditor de Calidad Lingüística (Quality Assurance). Tu tarea es revisar la traducción generada en el paso anterior (formato JSON) y certificar un 95% de precisión.
Criterios de revisión:
- Corrige cualquier error ortográfico, gramatical o de puntuación que se le haya escapado al traductor automático.
- Asegúrate de que no existan palabras sin sentido, texto residual en inglés o traducciones literales absurdas.
- Garantiza que el tono narrativo sea homogéneo en todo el documento.
Output: Entrega únicamente el JSON final definitivo perfectamente pulido y limpio, manteniendo la estructura (seq, text), listo para el usuario.
Texto a revisar: {json_nodo2}"""
                                
                                _logger.info("Página %s: Ejecutando NODO 3 (QA y Auditoría)...", page_num + 1)
                                resp3 = model.generate_content(prompt_nodo3, generation_config={"response_mime_type": "application/json"})
                                json_nodo3 = resp3.text.replace("```json", "").replace("```", "").strip()
                                
                                # Procesar respuesta final
                                final_list = json.loads(json_nodo3)
                                # Gemini puede devolver una lista de diccionarios o un diccionario directo
                                if isinstance(final_list, list):
                                    for item in final_list:
                                        translations_map[str(item.get("seq", ""))] = item.get("text", "")
                                        translations_map[item.get("seq")] = item.get("text", "")
                                elif isinstance(final_list, dict):
                                    translations_map = final_list
                                
                                _logger.info("Página %s: Pipeline Multi-Agente procesado con éxito.", page_num + 1)
                            except Exception as gemini_err:
                                _logger.warning("Error en el Pipeline Multi-Agente para la página %s (reintentando con fallback individual): %s", page_num, gemini_err)
                                translations_map = {}

                        # Fallback individual para rellenar vacíos o si falló el lote completo (optimizado en lote/batch)
                        trans = None
                        missing_blocks = [b for b in blocks_to_translate if str(b["seq"]) not in translations_map and b["seq"] not in translations_map]
                        if missing_blocks:
                            if not trans:
                                # Extracción dinámica del código ISO de 2 letras para deep_translator
                                dt_source = 'auto' if self.source_lang == 'auto' else self.source_lang.split('_')[0]
                                dt_target = self.target_lang.split('_')[0]
                                
                                # Ajuste técnico: Deep Translator suele procesar el Chino como 'zh-CN'
                                if dt_source == 'zh': dt_source = 'zh-CN'
                                if dt_target == 'zh': dt_target = 'zh-CN'

                                trans = GoogleTranslator(source=dt_source, target=dt_target)
                            
                            delimiter = " [para] "
                            # Unir los textos de los bloques faltantes con el delimitador
                            combined_text = delimiter.join(b["text"] for b in missing_blocks)
                            try:
                                translated_combined = trans.translate(combined_text)
                                # Separar el resultado por el delimitador
                                translated_parts = translated_combined.split(" [para] ")
                                if len(translated_parts) == len(missing_blocks):
                                    for i, b in enumerate(missing_blocks):
                                        translations_map[str(b["seq"])] = translated_parts[i].strip()
                                        translations_map[b["seq"]] = translated_parts[i].strip()
                                else:
                                    _logger.warning("Página %s: Discrepancia en división de traducción fallback (esperados %s, obtenidos %s). Reintentando individualmente.", 
                                                     page_num + 1, len(missing_blocks), len(translated_parts))
                                    for b in missing_blocks:
                                        try:
                                            res = trans.translate(b["text"])
                                            translations_map[str(b["seq"])] = res
                                            translations_map[b["seq"]] = res
                                        except Exception as single_err:
                                            _logger.warning("Error fallback individual en página %s, bloque %s: %s", page_num + 1, b["seq"], single_err)
                                            translations_map[str(b["seq"])] = b["text"]
                                            translations_map[b["seq"]] = b["text"]
                            except Exception as trans_err:
                                _logger.warning("Error fallback consolidada en página %s: %s. Reintentando individualmente.", page_num + 1, trans_err)
                                for b in missing_blocks:
                                    try:
                                        res = trans.translate(b["text"])
                                        translations_map[str(b["seq"])] = res
                                        translations_map[b["seq"]] = res
                                    except Exception as single_err:
                                        translations_map[str(b["seq"])] = b["text"]
                                        translations_map[b["seq"]] = b["text"]

                        # Guardado incremental de los segmentos traducidos
                        for b in blocks_to_translate:
                            seq_val = b["seq"]
                            translated = translations_map.get(str(seq_val)) or translations_map.get(seq_val) or b["text"]
                            self.env['translation.job.line'].create({
                                'job_id': self.id,
                                'page_number': page_num,
                                'sequence': seq_val,
                                'text_original': b["text"],
                                'text_translated': translated,
                                'is_heading': b.get("is_heading", False),
                                'alignment': b.get("alignment", "left"),
                            })
                    
                    # FUERA DEL SAVEPOINT: Realizar commit incremental seguro
                    self._safe_commit()
                    _logger.info("Página %s: Guardado completado en base de datos.", page_num + 1)

                except Exception as page_err:
                    _logger.exception("Error crítico controlado al procesar la página %s:", page_num)
                    continue

            # 4. Generación final del archivo DOCX
            _logger.info("Iniciando compilación del archivo DOCX final...")
            doc_word = docx.Document()
            lines = self.env['translation.job.line'].search([('job_id', '=', self.id)], order='page_number, sequence')
            
            if not lines:
                _logger.warning("No se encontraron líneas traducidas. Se creará un documento vacío.")
                doc_word.add_paragraph("No se extrajo ni tradujo ningún contenido del PDF original.")
            else:
                for line in lines:
                    # Crear elemento (título o párrafo)
                    if line.is_heading:
                        p = doc_word.add_heading(line.text_translated, level=1)
                    else:
                        p = doc_word.add_paragraph(line.text_translated)
                    
                    # Aplicar alineación
                    if line.alignment == 'center':
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif line.alignment == 'right':
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
            buffer = io.BytesIO()
            doc_word.save(buffer)
            
            # Reintentar el guardado final en caso de errores de concurrencia/serialización de base de datos
            max_retries = 3
            for attempt in range(max_retries):
                try:
                    with self.env.cr.savepoint():
                        self.write({
                            'docx_file': base64.b64encode(buffer.getvalue()),
                            'docx_filename': f"{self.name}_final.docx",
                            'state': 'done'
                        })
                    self._safe_commit()
                    _logger.info("Procesamiento de traducción finalizado con éxito (Intento %s/%s).", attempt + 1, max_retries)
                    break
                except Exception as write_err:
                    if attempt < max_retries - 1:
                        _logger.warning("Error de concurrencia al guardar resultado final (intento %s/%s), reintentando...: %s", 
                                       attempt + 1, max_retries, write_err)
                        time.sleep(1)
                    else:
                        raise write_err

        except Exception as e:
            _logger.exception("Error crítico no recuperable en _process_translation_background:")
            try:
                # Reintentar marcar como error para asegurar que se libere el estado visual
                for attempt in range(3):
                    try:
                        with self.env.cr.savepoint():
                            if not self._is_test_mode():
                                self.env.cr.rollback()
                            self.write({'state': 'error'})
                        self._safe_commit()
                        _logger.info("Estado del trabajo actualizado a 'error' en base de datos.")
                        break
                    except Exception as write_err:
                        if attempt < 2:
                            time.sleep(1)
                        else:
                            _logger.error(f"No se pudo establecer el estado de error en base de datos: {write_err}")
            except Exception as outer_err:
                _logger.error(f"Fallo crítico al intentar escribir el estado de error: {outer_err}")
        finally:
            # Eliminar del set global de hilos activos
            with self._active_jobs_lock:
                self._active_jobs.discard(self.id)
                
            if not self._is_test_mode():
                try:
                    self.env.cr.close()
                    _logger.info("Cursor de base de datos de fondo cerrado correctamente.")
                except Exception as close_err:
                    _logger.error(f"Fallo al cerrar el cursor de base de datos de fondo: {close_err}")